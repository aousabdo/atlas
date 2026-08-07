"""Original requirements -> today's systems, read from the workbook sheet
original_to_current_crosswalk.

Ported from build_cuas_tool_v4.py:944-977. This used to be a hardcoded literal
in the HTML template that the build never injected, so it could not follow the
source data. Never reintroduce a fallback copy: a stale crosswalk silently
falsifies the requirement-attrition metric.
"""
import openpyxl

KNOWN_STATUSES = {
    "Split out", "Renamed / split out", "Partly carried forward",
    "Condensed", "Didn't keep",
}

_NO_EQUIVALENT = "no matching system row"


class CrosswalkError(ValueError):
    """The workbook has no usable crosswalk sheet."""


def _normalise_status(raw):
    """Fold a spelled-out dropped status to the short label the UI shows.

    Real matrices write this status as a sentence rather than a token, and with
    a curly apostrophe, so the substring test below is deliberately loose. The
    fragments it matches have to stay literal for the parser to work on a real
    workbook, which is why the data guard allows them by name.
    """
    status = str(raw or "").strip() or "Unknown"
    lowered = status.lower().replace("’", "'")
    if "didn't keep" in lowered or "did not keep" in lowered:
        return "Didn't keep"
    return status


def _split_current(cell):
    if not cell or _NO_EQUIVALENT in str(cell).lower():
        return []
    return [c.strip() for c in str(cell).split(";") if c.strip()]


def read_crosswalk(xlsx_path):
    """-> [{orig, sys, current: [str], status}] in sheet order."""
    try:
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    except Exception as exc:
        raise CrosswalkError(f"could not open workbook: {exc}") from exc

    try:
        sheet = next((s for s in wb.sheetnames if "crosswalk" in s.lower()), None)
        if not sheet:
            raise CrosswalkError(
                "no 'original_to_current_crosswalk' sheet; the requirement "
                "attrition metric has no source"
            )
        rows = []
        for i, row in enumerate(wb[sheet].iter_rows(values_only=True)):
            if i == 0 or not row or not row[0]:
                continue
            rows.append({
                "orig": str(row[0]).strip(),
                "sys": str(row[1] or "").strip() if len(row) > 1 else "",
                "current": _split_current(row[2] if len(row) > 2 else None),
                "status": _normalise_status(row[3] if len(row) > 3 else None),
            })
        if not rows:
            raise CrosswalkError("crosswalk sheet is empty")
        return rows
    finally:
        wb.close()
