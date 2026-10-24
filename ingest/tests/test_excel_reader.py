"""How the Matrix reader opens a workbook and finds its integrations column.

The integrations cases, from the list check down to the dropped spelling, are
mirrored case for case in src/data/__tests__/localfile.test.ts, in "the
integrations column is found under every spelling the ingest accepts". The browser and this ingest must read
the same file the same way: a header only one of them recognises gives the CLI
bundle no mined links and the browser a full set, from identical bytes. Change
a case here and change it there.

The cells-that-are-not-text cases are mirrored the same way, in "a header or
name cell that is not text is skipped as the ingest skips it".
"""
import re
import zipfile

import openpyxl
import pytest

from atlas_ingest import excel
from atlas_ingest.excel import INTEGRATION_HEADERS, read_excel
from atlas_ingest.links import extract_links

PROSE = "Feeds Parity Gauge over the shared bus."

# Every spelling the reader accepts, in the order it tries them. The column is
# put first on purpose: column A is index 0, which an `or` chain reads as absent.
INTEGRATION_CASES = [
    "Existing Interfaces",
    "Currrent Integrations",
    "Current Integrations",
]


def _matrix(tmp_path, header, rows, name="parity.xlsx"):
    """A Matrix sheet with a title and a blank row above the header, the way the
    reader expects to find one."""
    path = tmp_path / name
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Matrix"
    ws.append(["Parity fixture"])
    ws.append([])
    ws.append(header)
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def test_the_case_list_is_the_list_the_reader_uses():
    assert list(INTEGRATION_HEADERS) == INTEGRATION_CASES


@pytest.mark.parametrize("header", INTEGRATION_CASES)
def test_every_spelling_yields_the_prose_and_its_mined_link(tmp_path, header):
    path = _matrix(tmp_path, [header, "Project/System", "Owner Organization"], [
        [PROSE, "Parity Widget", "Fixture Org"],
        [None, "Parity Gauge", "Fixture Org"],
    ])
    widget, gauge = read_excel(path)

    assert widget["integ"] == PROSE
    assert gauge["integ"] == ""
    assert [(link["from"], link["to"]) for link in extract_links([widget, gauge])] == [
        (widget["id"], gauge["id"]),
    ]


def test_the_first_spelling_in_the_list_wins_wherever_its_column_sits(tmp_path):
    path = _matrix(tmp_path, [
        "Current Integrations", "Project/System", "Currrent Integrations",
        "Existing Interfaces",
    ], [["from current", "Parity Widget", "from currrent", "from existing"]])
    assert [s["integ"] for s in read_excel(path)] == ["from existing"]


def test_the_doubled_letter_outranks_the_plain_spelling(tmp_path):
    path = _matrix(tmp_path, [
        "Current Integrations", "Project/System", "Currrent Integrations",
    ], [["from current", "Parity Widget", "from currrent"]])
    assert [s["integ"] for s in read_excel(path)] == ["from currrent"]


def test_a_spelling_outside_the_list_is_dropped_and_mines_nothing(tmp_path):
    path = _matrix(tmp_path, ["Integrations", "Project/System", "Owner Organization"], [
        [PROSE, "Parity Widget", "Fixture Org"],
        [None, "Parity Gauge", "Fixture Org"],
    ])
    systems = read_excel(path)
    assert [s["integ"] for s in systems] == ["", ""]
    assert extract_links(systems) == []


# --- read-only, as the module docstring says ----------------------------------


def test_the_workbook_opens_read_only(tmp_path, monkeypatch):
    """The module docstring promised read_only for a long time while the call
    did not pass it, and nothing noticed."""
    seen = {}
    real = excel.openpyxl.load_workbook

    def spy(*args, **kwargs):
        seen.update(kwargs)
        return real(*args, **kwargs)

    monkeypatch.setattr(excel.openpyxl, "load_workbook", spy)
    read_excel(_matrix(tmp_path, ["Project/System"], [["Parity Widget"]]))
    assert seen.get("read_only") is True


def test_a_short_row_in_a_sheet_with_no_dimension_reads_as_empty_cells(tmp_path):
    """Read-only is what makes this case exist. With no declared dimension,
    openpyxl ends each row at its last filled cell, so a row can stop before the
    integrations or Risk Level column. Indexing past it raised IndexError; the
    browser reads those cells as empty, and so must this."""
    full = _matrix(tmp_path, [
        "Project/System", "Owner Organization", "Existing Interfaces", "Risk Level",
    ], [
        ["Parity Widget", "Fixture Org", PROSE, "high"],
        ["Parity Gauge"],
    ])
    bare = tmp_path / "no-dimension.xlsx"
    with zipfile.ZipFile(full) as src, zipfile.ZipFile(bare, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith("xl/worksheets/"):
                data, count = re.subn(rb"<dimension [^>]*/>", b"", data)
                assert count == 1, "the fixture no longer declares a dimension to strip"
            out.writestr(item, data)

    widget, gauge = read_excel(bare)
    assert (widget["integ"], widget["risk"], widget["risk_source"]) == (PROSE, "high", "explicit")
    assert (gauge["integ"], gauge["risk_source"]) == ("", "inferred")


# --- cells that are not text ---------------------------------------------------


def test_a_number_in_the_header_row_is_skipped_and_the_next_column_still_read(tmp_path):
    """A year typed above a column is a number, not a header. The browser maps
    only cells that are text after trimming; this reader called .strip() on the
    number and raised AttributeError, so the CLI refused a file the browser
    loaded."""
    path = _matrix(tmp_path, [
        "Project/System", 2026, "   ", "Owner Organization", "Existing Interfaces",
    ], [["Parity Widget", "under the year", "under the blank", "Fixture Org", PROSE]])
    [widget] = read_excel(path)
    assert (widget["name"], widget["integ"]) == ("Parity Widget", PROSE)
    assert "Owner: Fixture Org" in widget["detail"]


def test_a_row_whose_name_is_only_whitespace_is_skipped(tmp_path):
    """The browser tests the trimmed name; this reader tested the raw cell, so
    a name of spaces became a system with an empty name in the CLI bundle
    only."""
    path = _matrix(tmp_path, ["Project/System", "Owner Organization"], [
        ["Parity Widget", "Fixture Org"],
        ["   ", "Fixture Org"],
        ["Parity Gauge", "Fixture Org"],
    ])
    assert [s["name"] for s in read_excel(path)] == ["Parity Widget", "Parity Gauge"]
